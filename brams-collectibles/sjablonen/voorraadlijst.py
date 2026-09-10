#!/usr/bin/env python3
"""Invullijst voor Bram: welk artikel hoort bij welke fotogroep, en wat kost het."""
import pathlib
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, GOUD, ROOM = '0D1B2A', 'C9A24B', 'F2EFE6'
INVUL = PatternFill('solid', fgColor='FFF7DC')      # geel = hier invullen
LICHT = PatternFill('solid', fgColor='F4F2EC')
dun = Side(style='thin', color='C9C4B8')
RAND = Border(left=dun, right=dun, top=dun, bottom=dun)

# Wat er per groep al bekend is. 'zeker' betekent: staat al in de winkel.
GROEPEN = [
    (1,  '3696–3698',  3,  'Destined Rivals — achterkant', '', False),
    (2,  '3699–3703',  5,  'Chaos Rising Elite Trainer Box', 'BC-CR-ETB', True),
    (3,  '3704–3708',  5,  'Chaos Rising Booster Bundle', 'BC-CR-BND', True),
    (4,  '3709–3711',  3,  'Mega Greninja ex Premium Collection', 'BC-MEV-GREN', True),
    (5,  '3712–3716',  5,  'Chaos Rising, donkere doos', '', False),
    (6,  '3717–3719',  3,  'Prismatic Evolutions Super Premium Collection', '', False),
    (7,  '3720–3722',  3,  'Prismatic Evolutions Super Premium Collection', '', False),
    (8,  '3723–3732', 10,  'Ascended Heroes, klein pakje', '', False),
    (9,  '3733–3740',  8,  'achterkant van een doos', '', False),
    (10, '3741–3745',  5,  'Elite Trainer Box met zilveren legendary', '', False),
    (11, '3758–3768', 11,  'Scarlet & Violet 151 Booster Bundle', 'BC-151-BND', True),
    (12, '3769–3785', 17,  'kartonnen doos met label', '', False),
    (13, '3786–3792',  7,  'Ascended Heroes Booster Bundle', '', False),
    (14, '3793–3794',  2,  'verzenddoos', '', False),
    (15, '3795–3796',  2,  'verzenddoos, rood', '', False),
]

KOPPEN = ['Groep', 'Foto', "Foto's", 'Wat ik erin zie', 'Zeker?',
          'Naam van het artikel', 'Aantal', 'Inkoop per stuk', 'Verkoopprijs', 'Opmerking']
BREED = [7, 22, 8, 34, 9, 40, 9, 15, 14, 30]

wb = Workbook()
ws = wb.active
ws.title = 'Invullijst'

ws['A1'] = 'Brams Collectibles — voorraadlijst'
ws['A1'].font = Font(name='Arial', size=16, bold=True, color=NAVY)
ws['A2'] = ('Vul de gele kolommen in. De foto in kolom B is een van de opnames uit die groep, '
            'zodat je ziet welk artikel bedoeld wordt.')
ws['A2'].font = Font(name='Arial', size=10, color='5A6A78')
ws.merge_cells('A2:J2')

ws['A4'] = 'Wat er van je gevraagd wordt'
ws['A4'].font = Font(name='Arial', size=11, bold=True, color=NAVY)
uitleg = [
    'Naam van het artikel — alleen waar "Zeker?" leeg is, of waar mijn omschrijving niet klopt.',
    'Aantal — hoeveel je er van dit artikel hebt. Niet het aantal foto\'s.',
    'Inkoop per stuk — wat je ervoor betaald hebt. Zonder dit getal weten we niet of er winst in zit.',
    'Verkoopprijs — wat je wilt vragen. Leeg laten mag; dan kijken we samen naar Cardmarket.',
]
for i, r in enumerate(uitleg):
    c = ws.cell(row=5 + i, column=1, value='• ' + r)
    c.font = Font(name='Arial', size=10)
    ws.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=10)

KOP = 10
for k, (naam, breedte) in enumerate(zip(KOPPEN, BREED), start=1):
    c = ws.cell(row=KOP, column=k, value=naam)
    c.font = Font(name='Arial', size=10, bold=True, color=ROOM)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = RAND
    ws.column_dimensions[get_column_letter(k)].width = breedte
ws.row_dimensions[KOP].height = 30

for n, (groep, reeks, aantal_fotos, zien, sku, zeker) in enumerate(GROEPEN):
    r = KOP + 1 + n
    ws.row_dimensions[r].height = 92
    waarden = [groep, '', aantal_fotos, zien, 'staat al in de winkel' if zeker else '',
               sku and zien or '', '', '', '', '']
    for k, v in enumerate(waarden, start=1):
        c = ws.cell(row=r, column=k, value=v if v != '' else None)
        c.font = Font(name='Arial', size=10)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = RAND
        if k in (6, 7, 8, 9, 10):
            c.fill = INVUL
        elif not zeker:
            c.fill = LICHT
    ws.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=r, column=8).number_format = '€ #,##0.00'
    ws.cell(row=r, column=9).number_format = '€ #,##0.00'
    ws.cell(row=r, column=7).alignment = Alignment(horizontal='center', vertical='center')

    mini = pathlib.Path(f'/tmp/mini/g{groep:02d}.png')
    if mini.exists():
        img = XLImage(str(mini))
        img.height, img.width = 115, int(115 * img.width / img.height)
        ws.add_image(img, f'B{r}')

# voorbeeldregel, zodat het formaat duidelijk is
vb = KOP + len(GROEPEN) + 2
ws.cell(row=vb, column=1, value='Voorbeeld').font = Font(name='Arial', size=10, bold=True, italic=True)
for k, v in [(6, 'Prismatic Evolutions Pokémon Center ETB'), (7, 3), (8, 410), (9, 599),
             (10, 'twee met acrylhoes, één zonder')]:
    c = ws.cell(row=vb, column=k, value=v)
    c.font = Font(name='Arial', size=10, italic=True, color='7A7A7A')
    c.border = RAND
ws.cell(row=vb, column=8).number_format = '€ #,##0.00'
ws.cell(row=vb, column=9).number_format = '€ #,##0.00'

slot = vb + 2
ws.cell(row=slot, column=1, value=(
    'Groep 6 en 7 lijken hetzelfde artikel van twee kanten. Is dat zo, zet dan bij groep 7 '
    '"zelfde als 6". Groep 12, 14 en 15 zijn kartonnen dozen — is dat een verzegelde case met '
    'meerdere stuks erin, of verpakkingsmateriaal?')).font = Font(name='Arial', size=10, italic=True, color='A9532F')
ws.merge_cells(start_row=slot, start_column=1, end_row=slot, end_column=10)

ws.freeze_panes = f'A{KOP + 1}'
ws.sheet_view.showGridLines = False

doel = '/home/user/vloerenoffertepro/brams-collectibles/voorraadlijst-invullen.xlsx'
wb.save(doel)
print('geschreven:', doel)
